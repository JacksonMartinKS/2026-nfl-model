#!/usr/bin/env python3
"""Rebuild forecast-data.js for the NFL forecast site from your two workbooks.

    python update_data.py "NFL SIMULATOR RESULTS.xlsx" NFL_SPREADSHEET_v2.xlsx

Options:
    --season 2026            season to read (default 2026)
    --out forecast-data.js   output path (default: next to this script)

What it reads
  Simulator workbook
    * "<season>_sim - pre_summary", "<season>_sim - wk1_summary", ... one sheet per
      week, with the columns your sim writes: team, mean_wins, stdev_wins, poff,
      p_div, p_conf_win, p_sb, exp_seed, p_seed1 ... p_seed7
    * "<yy> Personal" for team names, conference and division
  Elo workbook
    * "<yy>" season sheet: game rows under "<season> Week N" labels with
      away, home, away score, home score in columns B-E (used for W-L-T records),
      plus the pregame columns (Away/Home Win Chance, Predicted Home Team Victory
      Margin, Game Quality Index, Game ELOs, Date, rest days, Stadium Coord., and the
      rest / field / QB / travel adjustments) for the Game forecasts view, every week
      through the upcoming one, with final scores once games are played

Every week sheet that exists is included, so after Week 6 runs and its summary sheet
is in the workbook, re-run this and push index.html + forecast-data.js.
Needs openpyxl (pip install openpyxl). Save the workbooks in Excel first so cached
formula values are present.
"""
import argparse
import datetime as dt
import json
import math
import re
from pathlib import Path

from openpyxl import load_workbook

FIELDS = ["mean_wins", "stdev_wins", "poff", "p_div", "p_conf_win", "p_sb"]


def sheet_rows(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]


def team_meta(rows):
    teams, seen = [], set()
    for r in rows:
        if len(r) < 5:
            continue
        name, abbr, conf, div = r[0], r[1], r[3], r[4]
        if isinstance(abbr, str) and re.fullmatch(r"[A-Z]{2,3}", abbr) and conf in ("AFC", "NFC") and abbr not in seen:
            teams.append({"name": str(name).strip(), "abbr": abbr, "conf": conf, "div": str(div).strip()})
            seen.add(abbr)
    if len(teams) != 32:
        raise SystemExit(f"Expected 32 teams in the Personal sheet, found {len(teams)}")
    return teams


def week_summary(rows, abbrs, sheet_name):
    h = next((i for i, r in enumerate(rows) if r and str(r[0]).strip().lower() == "team"), None)
    if h is None:
        raise SystemExit(f'No "team" header row in {sheet_name}')
    hdr = [str(c).strip() if c is not None else "" for c in rows[h]]
    out = {}
    for r in rows[h + 1:]:
        if not r or r[0] not in abbrs:
            continue
        d = dict(zip(hdr, r))
        out[r[0]] = [round(float(d[k]), 5) for k in FIELDS] + [[round(float(d[f"p_seed{s}"]), 5) for s in range(1, 8)]]
    missing = abbrs - set(out)
    if missing:
        raise SystemExit(f"{sheet_name} is missing teams: {sorted(missing)}")
    return out


def records_by_week(rows, abbrs, season, last_week):
    """Cumulative [W, L, T] after each week 0..last_week from final scores."""
    games, wk = {}, None
    for r in rows:
        if not r:
            continue
        label = str(r[0] or "").strip()
        if label:
            m = re.match(rf"{season} Week (\d+)", label)
            wk = int(m.group(1)) if m else None
        if wk is None or wk > last_week or len(r) < 5:
            continue
        away, home = r[1], r[2]
        if away not in abbrs or home not in abbrs:
            continue
        try:
            sa, sh = float(r[3]), float(r[4])
        except (TypeError, ValueError):
            continue  # unplayed game
        games.setdefault(wk, []).append((away, home, sa, sh))
    rec = {a: [0, 0, 0] for a in abbrs}
    out = {0: {a: v[:] for a, v in rec.items()}}
    for w in range(1, last_week + 1):
        for away, home, sa, sh in games.get(w, []):
            if sa > sh:
                rec[away][0] += 1; rec[home][1] += 1
            elif sh > sa:
                rec[home][0] += 1; rec[away][1] += 1
            else:
                rec[away][2] += 1; rec[home][2] += 1
        out[w] = {a: v[:] for a, v in rec.items()}
    return out, {w: len(g) for w, g in games.items()}


INTL_SITES = [
    ("London", 51.51, -0.13), ("Dublin", 53.34, -6.26), ("Munich", 48.22, 11.62), ("Berlin", 52.51, 13.24),
    ("Frankfurt", 50.07, 8.65), ("Paris", 48.84, 2.25), ("Madrid", 40.45, -3.69), ("Mexico City", 19.30, -99.15),
    ("São Paulo", -23.55, -46.64), ("Rio de Janeiro", -22.91, -43.23), ("Melbourne", -37.82, 144.98),
    ("Toronto", 43.64, -79.39),
]


def site_name(coord):
    """Name an international venue from the 'lat, lon' stadium coordinate, else None."""
    try:
        lat, lon = (float(x) for x in str(coord).split(","))
    except ValueError:
        return None
    def dist(s):
        return math.hypot(s[1] - lat, (s[2] - lon) * math.cos(math.radians(lat)))
    best = min(INTL_SITES, key=dist)
    if dist(best) < 2.5:
        return best[0]
    if not (15 < lat < 62 and -170 < lon < -50):
        return "International"
    return None


def excel_date(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    try:
        return (dt.date(1899, 12, 30) + dt.timedelta(days=int(float(v)))).isoformat()
    except (TypeError, ValueError):
        return None


def season_games(rows, abbrs, season):
    """Pregame Elo detail (plus the final score once played) for every game, keyed by week."""
    hi = next((i for i, r in enumerate(rows) if r and "Away Win Chance" in [str(c).strip() for c in r if c is not None]), None)
    if hi is None:
        print("No 'Away Win Chance' header row found; skipping game forecasts")
        return {}
    hdr = [str(c).strip() if c is not None else "" for c in rows[hi]]
    col = lambda name, start=0: hdr.index(name, start)
    c = {
        "pa": col("Away Win Chance"), "m": col("Predicted Home Team Victory Margin"), "q": col("Game Quality Index"),
        "ae": col("Away Game ELO"), "he": col("Home Game ELO"), "d": col("Date"), "site": col("Stadium Coord."),
        "ard": col("Away Rest Days"), "hrd": col("Home Rest Days"),
        "ar": col("Away Rest ELO Adj."), "af": col("Away Field Adv."), "at": col("Away Travel ELO Adj."),
        "hr": col("Home Rest ELO Adj."), "hf": col("Home Field Adv."), "ht": col("Home Travel ELO Adj."),
    }
    c["aq"] = col("QB Mod.", c["af"])
    c["hq"] = col("QB Mod.", c["hf"])
    last_col = max(c.values())
    by_week, wk = {}, None
    for r in rows:
        if not r:
            continue
        label = str(r[0] or "").strip()
        if label:
            m = re.match(rf"{season} Week (\d+)", label)
            wk = int(m.group(1)) if m else None
        if wk is None or len(r) <= last_col:
            continue
        a, h = r[1], r[2]
        if a not in abbrs or h not in abbrs:
            continue
        try:
            f = lambda k, nd=2: round(float(r[c[k]]), nd)
            g = {
                "d": excel_date(r[c["d"]]), "a": a, "h": h,
                "pa": f("pa", 4), "m": f("m"), "q": f("q", 1),
                # [game Elo, rest days, rest adj, field adv, QB adj, travel adj]
                "A": [f("ae", 1), int(float(r[c["ard"]])), f("ar"), f("af"), f("aq"), f("at")],
                "H": [f("he", 1), int(float(r[c["hrd"]])), f("hr"), f("hf"), f("hq"), f("ht")],
                "site": site_name(r[c["site"]]),
            }
        except (TypeError, ValueError):
            continue
        try:
            g["s"] = [int(float(r[3])), int(float(r[4]))]  # final score, away-home
        except (TypeError, ValueError):
            pass
        by_week.setdefault(wk, []).append(g)
    full_hfa = max((g["H"][3] for gl in by_week.values() for g in gl), default=0)
    for gl in by_week.values():
        for g in gl:
            if not g["site"] and g["H"][3] < full_hfa:
                g["site"] = "Neutral site"
    return by_week


def next_game_elo(by_week, team, after_week):
    """[base Elo, QB adj] going into the team's first game after `after_week` (skips byes)."""
    for w in sorted(k for k in by_week if k > after_week):
        for g in by_week[w]:
            side = "A" if g["a"] == team else "H" if g["h"] == team else None
            if side:
                s = g[side]
                return [round(s[0] - (s[2] + s[3] + s[4] + s[5]), 1), round(s[4], 1)]
    return [None, None]


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("sim_xlsx")
    ap.add_argument("elo_xlsx")
    ap.add_argument("--season", type=int, default=2026)
    ap.add_argument("--out", default=str(Path(__file__).with_name("forecast-data.js")))
    args = ap.parse_args()
    season, yy = args.season, str(args.season)[2:]

    sim = load_workbook(args.sim_xlsx, read_only=True, data_only=True)
    teams = team_meta(sheet_rows(sim[f"{yy} Personal"]))
    abbrs = {t["abbr"] for t in teams}

    weeks = []
    for key in ["pre"] + [f"wk{n}" for n in range(1, 19)]:
        name = f"{season}_sim - {key}_summary"
        if name not in sim.sheetnames:
            continue
        rows = week_summary(sheet_rows(sim[name]), abbrs, name)
        label = "Preseason" if key == "pre" else f"Week {key[2:]}"
        weeks.append({"key": key, "label": label, "teams": rows})
        print(f"{name}: wins sum {sum(v[0] for v in rows.values()):.1f}, playoff sum {sum(v[2] for v in rows.values()):.2f}")
    if not weeks:
        raise SystemExit(f'No "{season}_sim - ..._summary" sheets found')

    last_week = 0 if weeks[-1]["key"] == "pre" else int(weeks[-1]["key"][2:])
    elo = load_workbook(args.elo_xlsx, read_only=True, data_only=True)
    elo_rows = sheet_rows(elo[yy])
    recs, per_week = records_by_week(elo_rows, abbrs, season, last_week)
    print("games with scores per week:", per_week)
    by_week = season_games(elo_rows, abbrs, season)
    games = []
    for w in range(1, min(last_week + 1, 18) + 1):
        if by_week.get(w):
            played = {t for g in by_week[w] for t in (g["a"], g["h"])}
            games.append({"week": w, "label": f"Week {w}", "list": by_week[w], "byes": sorted(abbrs - played)})
    print("game forecast weeks:", [g["week"] for g in games])

    for wk in weeks:
        n = 0 if wk["key"] == "pre" else int(wk["key"][2:])
        for a, row in wk["teams"].items():
            row.append(recs.get(n, recs[0])[a])
            row.extend(next_game_elo(by_week, a, n))  # [base Elo, QB adj] for the next game

    data = {
        "season": season,
        "updated": dt.date.today().isoformat(),
        "fields": FIELDS + ["p_seed1..7", "record_wlt", "next_base_elo", "next_qb_adj"],
        "teams": teams,
        "weeks": weeks,
        "games": games,
    }
    Path(args.out).write_text(
        "/* Jackson Martin Forecasts · NFL forecast data. Generated by update_data.py; do not edit by hand. */\n"
        "window.NFL_DATA = " + json.dumps(data, separators=(",", ":")) + ";\n",
        encoding="utf-8",
    )
    print(f"Wrote {args.out} ({len(weeks)} weeks, through {weeks[-1]['label']})")


if __name__ == "__main__":
    main()
